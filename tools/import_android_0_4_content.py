#!/usr/bin/env python3
"""Deterministic content importer for DUBL Android 0.4.

Currently imports the approved martial-arts scope from the Master of Melee
rulebook into development_catalog.json. Archmage spell import is handled by
--archmage once that source is supplied.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Iterator, Sequence

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


MARTIAL_SECTIONS = {
    "Рукопашные искусства": "Рукопашные",
    "Оружейные искусства": "Оружейные",
}
STOP_MARTIAL_SECTION = "Гибридные искусства"
ANDROID_MAGIC_SCHOOLS = {
    "Боевая магия",
    "Воплощение",
    "Друид",
    "Магия крови",
    "Молитва",
    "Некромантия",
    "Ограждение",
    "Призыв",
    "Природа",
    "Прорицание",
    "Разрушение",
    "Разум",
    "Трансмутация",
}
MAGIC_SCHOOL_ALIASES = {
    "молитвы": "Молитва",
    "молитва": "Молитва",
}
STYLE_NAME_ALIASES = {
    "карате": "Каратэ",
    "ган ката": "Ган-Ката",
}


def clean(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("\u00a0", " ")).strip()


def normalize(value: str) -> str:
    value = clean(value).lower().replace("ё", "е")
    value = value.replace("–", "-").replace("—", "-")
    return value.strip(" .,:;!?")


def stable_id(kind: str, name: str) -> str:
    digest = hashlib.sha1(f"{kind}:{normalize(name)}".encode("utf-8")).hexdigest()[:16]
    return f"martial_{kind}_{digest}"


def canonical_style_name(title: str) -> str:
    alias = STYLE_NAME_ALIASES.get(normalize(title))
    return alias or clean(title)


def iter_document_blocks(document: DocumentType) -> Iterator[Paragraph | Table]:
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def paragraph_text(block: Paragraph | Table) -> str:
    return block.text if isinstance(block, Paragraph) else ""


def is_blank(block: Paragraph | Table) -> bool:
    return isinstance(block, Paragraph) and not clean(block.text)


def first_int(pattern: str, text: str) -> int | None:
    match = re.search(pattern, text, flags=re.IGNORECASE)
    return int(match.group(1)) if match else None


def field_value(text: str, label: str, stop_labels: Sequence[str]) -> str:
    escaped_stops = "|".join(re.escape(label) for label in stop_labels)
    pattern = rf"(?:^|\n)\s*{re.escape(label)}\s*:\s*(.*?)(?=\n\s*(?:{escaped_stops})\s*:|\Z)"
    match = re.search(pattern, text, flags=re.IGNORECASE | re.DOTALL)
    if not match:
        return ""
    return clean(match.group(1))


def parse_cost_rank(text: str) -> tuple[int | None, int]:
    cost = first_int(r"Стоимость\s*:\s*(\d+)", text)
    if cost is None:
        cost = first_int(r"^\s*(\d+)\s*,", text)
    rank = first_int(r"Ранг(?:и|ов)?\s*:\s*(\d+)", text)
    if rank is None:
        rank = first_int(r",\s*(\d+)\s*ранг", text)
    return cost, rank or 1


def extract_requirement(text: str) -> str:
    return field_value(
        text,
        "Требование",
        ["Выгода", "Примечание", "Особое", "Ограничение", "Описание", "Стоимость", "Ранг", "Ранги"],
    ).strip(" .")


def extract_benefit(text: str) -> str:
    return field_value(
        text,
        "Выгода",
        ["Примечание", "Особое", "Ограничение", "Стоимость", "Требование", "Ранг", "Ранги"],
    ).strip()


def extract_prefixed(text: str, label: str) -> str:
    return field_value(
        text,
        label,
        ["Описание", "История", "Обязательное условие", "Стоимость", "Ранг", "Ранги", "Требование", "Выгода", "Примечание", "Особое", "Ограничение"],
    ).strip()


def combine_parts(*parts: str) -> str:
    return "\n".join(part for part in (clean(p) for p in parts) if part)


def martial_entry(
    *,
    entry_id: str,
    name: str,
    category: str,
    cost: int,
    ranks: int,
    requirements: str,
    benefit: str,
    notes: str,
    tags: Sequence[str],
) -> dict:
    return {
        "id": entry_id,
        "name": clean(name),
        "section": "Боевые искусства",
        "category": clean(category),
        "cost": max(0, cost),
        "costType": "xp",
        "ranks": max(1, ranks),
        "requirements": clean(requirements) or "-",
        "benefit": clean(benefit),
        "notes": clean(notes),
        "tags": list(dict.fromkeys(clean(tag) for tag in tags if clean(tag))),
    }


@dataclass
class TechniqueOccurrence:
    name: str
    cost: int
    ranks: int
    requirements: str
    benefit: str
    notes: str
    source_style: str
    martial_type: str
    stance: bool = False

    @property
    def score(self) -> int:
        return len(self.requirements) + len(self.benefit) * 2 + len(self.notes)


@dataclass
class TechniqueAggregate:
    best: TechniqueOccurrence
    occurrences: list[TechniqueOccurrence] = field(default_factory=list)

    def add(self, occurrence: TechniqueOccurrence) -> None:
        self.occurrences.append(occurrence)
        if occurrence.score > self.best.score:
            self.best = occurrence

    def to_entry(self) -> dict:
        sources = sorted({item.source_style for item in self.occurrences}, key=normalize)
        types = sorted({item.martial_type for item in self.occurrences}, key=normalize)
        best = self.best
        requirement = best.requirements
        generated_single = normalize(requirement) == normalize(f"Боевые искусства ({best.source_style})")
        if len(sources) > 1 and (not requirement or generated_single):
            requirement = "Боевые искусства: " + ", ".join(sources[:-1]) + (f" или {sources[-1]}" if len(sources) > 1 else sources[0])
        category = "Общие приёмы" if len(sources) > 1 else f"Приёмы · {sources[0]}"
        tags = ["Боевые искусства", "Приём", *types, *sources]
        if any(item.stance for item in self.occurrences):
            tags.append("Стойка")
        return martial_entry(
            entry_id=stable_id("technique", best.name),
            name=best.name,
            category=category,
            cost=best.cost,
            ranks=best.ranks,
            requirements=requirement,
            benefit=best.benefit,
            notes=best.notes,
            tags=tags,
        )


def parse_style_root(style_name: str, martial_type: str, blocks: Sequence[Paragraph | Table]) -> dict:
    root_blocks: list[Paragraph] = []
    for block in blocks:
        if isinstance(block, Table):
            break
        text = clean(block.text)
        if text and normalize(text) in {"приемы", "стойки", "совершенная способность"}:
            break
        if isinstance(block, Paragraph):
            root_blocks.append(block)

    lines = [block.text.strip() for block in root_blocks if clean(block.text)]
    joined = "\n".join(lines)
    cost, ranks = parse_cost_rank(joined)
    if cost is None:
        raise ValueError(f"No style cost found for {style_name}")

    requirement = extract_requirement(joined)
    mandatory = ""
    for line in lines:
        if normalize(line).startswith("обязательное условие:"):
            mandatory = clean(line.split(":", 1)[1])
            break
    if mandatory:
        requirement = "; ".join(part for part in (requirement, mandatory) if part)

    benefit = extract_benefit(joined)
    if not benefit:
        raise ValueError(f"No style benefit found for {style_name}")

    metadata_start = next((i for i, line in enumerate(lines) if re.search(r"Стоимость\s*:", line, re.I)), len(lines))
    flavor_lines = []
    style_norm = normalize(style_name)
    for line in lines[:metadata_start]:
        one = clean(line)
        if not one or normalize(one) == style_norm:
            continue
        if normalize(one).startswith("обязательное условие:"):
            continue
        flavor_lines.append(one)
    special = next((clean(line.split(":", 1)[1]) for line in lines if normalize(line).startswith("особое:") and ":" in line), "")
    notes = combine_parts("\n".join(flavor_lines), f"Особое: {special}" if special else "")

    return martial_entry(
        entry_id=stable_id("style", style_name),
        name=style_name,
        category=f"{martial_type} стили",
        cost=cost,
        ranks=ranks,
        requirements=requirement,
        benefit=benefit,
        notes=notes,
        tags=["Боевые искусства", "Боевой стиль", martial_type],
    )


def parse_paragraph_techniques(
    style_name: str,
    martial_type: str,
    blocks: Sequence[Paragraph | Table],
) -> list[TechniqueOccurrence]:
    occurrences: list[TechniqueOccurrence] = []
    active = False
    for idx, block in enumerate(blocks):
        if not isinstance(block, Paragraph):
            continue
        text = clean(block.text)
        ntext = normalize(text)
        if ntext == "приемы":
            active = True
            continue
        if block.style.name == "Heading 3":
            if ntext == "приемы":
                active = True
            elif ntext in {"совершенная способность"}:
                active = False
            # Table-only stance sections are handled separately.
            continue
        if not active or not re.search(r"Стоимость\s*:\s*\d+", block.text, re.I):
            continue

        start = idx - 1
        while start >= 0 and not is_blank(blocks[start]) and not isinstance(blocks[start], Table):
            p = blocks[start]
            if isinstance(p, Paragraph) and p.style.name in {"Title", "Heading 2", "Heading 3"}:
                break
            start -= 1
        start += 1

        end = idx + 1
        while end < len(blocks) and not is_blank(blocks[end]) and not isinstance(blocks[end], Table):
            p = blocks[end]
            if isinstance(p, Paragraph) and p.style.name in {"Title", "Heading 2", "Heading 3"}:
                break
            end += 1

        chunk_paragraphs = [b for b in blocks[start:end] if isinstance(b, Paragraph) and clean(b.text)]
        if not chunk_paragraphs:
            continue
        name = clean(chunk_paragraphs[0].text)
        if normalize(name) in {"приемы", "стойки"} or ":" in name:
            continue
        chunk = "\n".join(p.text.strip() for p in chunk_paragraphs)
        cost, ranks = parse_cost_rank(chunk)
        if cost is None:
            continue
        requirement = extract_requirement(chunk) or f"Боевые искусства ({style_name})"
        benefit = extract_benefit(chunk)
        if not benefit:
            continue

        cost_pos = next((i for i, p in enumerate(chunk_paragraphs) if re.search(r"Стоимость\s*:", p.text, re.I)), 1)
        pre_cost = [clean(p.text) for p in chunk_paragraphs[1:cost_pos]]
        notes_parts = [part for part in pre_cost if part and not (part.startswith("(") and part.endswith(")"))]
        note_field = extract_prefixed(chunk, "Примечание")
        notes = combine_parts(" ".join(notes_parts), f"Примечание: {note_field}" if note_field else "")
        stance = "[стойка]" in normalize(chunk)
        occurrences.append(
            TechniqueOccurrence(
                name=name,
                cost=cost,
                ranks=ranks,
                requirements=requirement,
                benefit=benefit,
                notes=notes,
                source_style=style_name,
                martial_type=martial_type,
                stance=stance,
            )
        )
    return occurrences


def parse_table_techniques(
    style_name: str,
    martial_type: str,
    blocks: Sequence[Paragraph | Table],
) -> list[TechniqueOccurrence]:
    occurrences: list[TechniqueOccurrence] = []
    subsection = ""
    for block in blocks:
        if isinstance(block, Paragraph):
            text = normalize(block.text)
            if block.style.name == "Heading 3" or text in {"стойки", "приемы"}:
                subsection = text
            if text == "совершенная способность":
                subsection = ""
            continue
        if subsection not in {"стойки", "приемы"}:
            continue
        rows = block.rows
        if not rows:
            continue
        for row in rows[1:]:
            cells = [clean(cell.text) for cell in row.cells]
            if len(cells) < 4:
                continue
            name, cost_text, effect = cells[1], cells[2], cells[3]
            if not name or not cost_text or not effect:
                continue
            cost, ranks = parse_cost_rank(cost_text)
            if cost is None:
                continue
            requirement = extract_requirement(effect) or f"Боевые искусства ({style_name})"
            benefit = extract_benefit(effect) or effect
            occurrences.append(
                TechniqueOccurrence(
                    name=name,
                    cost=cost,
                    ranks=ranks,
                    requirements=requirement,
                    benefit=benefit,
                    notes="",
                    source_style=style_name,
                    martial_type=martial_type,
                    stance=subsection == "стойки" or "[стойка]" in normalize(effect),
                )
            )
    return occurrences


def parse_martial_arts(source: Path) -> tuple[list[dict], list[dict]]:
    doc = Document(source)
    blocks = list(iter_document_blocks(doc))
    current_section: str | None = None
    section_type: str | None = None
    styles: list[tuple[str, str, list[Paragraph | Table]]] = []
    active_name: str | None = None
    active_type: str | None = None
    active_blocks: list[Paragraph | Table] = []

    def flush() -> None:
        nonlocal active_name, active_type, active_blocks
        if active_name and active_type:
            styles.append((active_name, active_type, active_blocks))
        active_name = None
        active_type = None
        active_blocks = []

    for block in blocks:
        if not isinstance(block, Paragraph):
            if active_name:
                active_blocks.append(block)
            continue
        text = clean(block.text)
        if block.style.name == "Title" and text in MARTIAL_SECTIONS:
            flush()
            current_section = text
            section_type = MARTIAL_SECTIONS[text]
            continue
        if block.style.name == "Title" and text == STOP_MARTIAL_SECTION:
            flush()
            current_section = None
            section_type = None
            break
        if current_section and block.style.name == "Title" and text:
            flush()
            active_name = canonical_style_name(text)
            active_type = section_type
            active_blocks = []
            continue
        if active_name:
            active_blocks.append(block)
    flush()

    style_entries: list[dict] = []
    technique_map: dict[str, TechniqueAggregate] = {}
    for style_name, martial_type, style_blocks in styles:
        style_entries.append(parse_style_root(style_name, martial_type, style_blocks))
        occurrences = parse_paragraph_techniques(style_name, martial_type, style_blocks)
        occurrences += parse_table_techniques(style_name, martial_type, style_blocks)
        for occurrence in occurrences:
            key = normalize(occurrence.name)
            aggregate = technique_map.get(key)
            if aggregate is None:
                technique_map[key] = TechniqueAggregate(best=occurrence, occurrences=[occurrence])
            else:
                aggregate.add(occurrence)

    technique_entries = [aggregate.to_entry() for _, aggregate in sorted(technique_map.items())]
    style_entries.sort(key=lambda item: (normalize(item["category"]), normalize(item["name"])))
    technique_entries.sort(key=lambda item: (normalize(item["category"]), normalize(item["name"])))
    return style_entries, technique_entries


def update_development_catalog(catalog_path: Path, style_entries: list[dict], technique_entries: list[dict]) -> None:
    root = json.loads(catalog_path.read_text(encoding="utf-8"))
    existing = [entry for entry in root["entries"] if normalize(entry.get("section", "")) != "боевые искусства"]
    imported = style_entries + technique_entries
    ids = [entry["id"] for entry in imported]
    if len(ids) != len(set(ids)):
        raise ValueError("Generated martial-art IDs are not unique")
    root["version"] = "0.12.0"
    root["entries"] = existing + imported
    catalog_path.write_text(json.dumps(root, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def canonical_magic_school(name: str) -> str:
    normalized = normalize(name)
    alias = MAGIC_SCHOOL_ALIASES.get(normalized)
    if alias:
        return alias
    return next((school for school in ANDROID_MAGIC_SCHOOLS if normalize(school) == normalized), clean(name))


def parse_magic_schools(raw: str) -> list[str]:
    parts = [canonical_magic_school(part) for part in re.split(r"\s*/\s*", clean(raw)) if clean(part)]
    return list(dict.fromkeys(parts))


def parse_archmage_additions(source: Path) -> list[dict]:
    """Parse supported extra spell definitions from the Archmage book.

    Most additions are numbered Heading 3 blocks. A small set of Battle Magic
    additions is formatted as Heading 5 with all metadata in the following paragraph.
    We accept either form only when that next paragraph is a complete definition;
    this avoids re-importing stock Heading 5 spells whose metadata is split across
    several paragraphs.
    """
    doc = Document(source)
    additions: list[dict] = []
    paragraphs = doc.paragraphs
    for index, paragraph in enumerate(paragraphs[:-1]):
        heading = clean(paragraph.text)
        style_name = paragraph.style.name
        numbered_addition = style_name == "Heading 3" and bool(re.match(r"^\d+\.\s+", heading))
        inline_heading_addition = style_name == "Heading 5"
        if not numbered_addition and not inline_heading_addition:
            continue
        body = paragraphs[index + 1].text.strip()
        if "Школа:" not in body or "Стоимость:" not in body or "Описание:" not in body:
            continue

        name = clean(re.sub(r"^\d+\.\s*", "", heading)) if numbered_addition else heading
        school_raw = field_value(body, "Школа", ["Стоимость", "Время сотворения", "Дальность", "Область", "Действие", "Длительность", "Описание", "Усиление"])
        schools = parse_magic_schools(school_raw)
        if not schools or any(school not in ANDROID_MAGIC_SCHOOLS for school in schools):
            continue

        cost = first_int(r"(?:^|\n)\s*Стоимость\s*:\s*(\d+)", body)
        if cost is None:
            raise ValueError(f"No mana cost for Archmage spell {name}")
        time = field_value(body, "Время сотворения", ["Дальность", "Область", "Действие", "Длительность", "Описание", "Усиление"])
        range_value = field_value(body, "Дальность", ["Область", "Действие", "Длительность", "Описание", "Усиление"])
        area = field_value(body, "Область", ["Действие", "Длительность", "Описание", "Усиление"])
        action = field_value(body, "Действие", ["Длительность", "Описание", "Усиление"])
        duration = field_value(body, "Длительность", ["Описание", "Усиление"])
        description = field_value(body, "Описание", ["Усиление"])
        enhancement = field_value(body, "Усиление", [])
        if not description:
            raise ValueError(f"No description for Archmage spell {name}")

        school_text = " / ".join(schools)
        digest = hashlib.sha1(f"{normalize(name)}:{normalize(school_text)}".encode("utf-8")).hexdigest()[:16]
        entry = {
            "id": f"archmage_{digest}",
            "name": name,
            "section": "Заклинания",
            "category": school_text,
            "cost": cost,
            "school": school_text,
            "manaText": str(cost),
            "time": clean(time),
            "range": clean(range_value),
            "area": clean(area),
            "action": clean(action),
            "duration": clean(duration),
            "description": clean(description),
            "enhancement": clean(enhancement),
            "incomplete": False,
            "source": "Книга Архимага",
        }
        additions.append(entry)

    ids = [entry["id"] for entry in additions]
    if len(ids) != len(set(ids)):
        raise ValueError("Generated Archmage spell IDs are not unique")
    return additions


def update_magic_catalog(catalog_path: Path, additions: list[dict]) -> int:
    root = json.loads(catalog_path.read_text(encoding="utf-8"))
    baseline = [entry for entry in root.get("spells", []) if not str(entry.get("id", "")).startswith("archmage_")]
    existing_names = {normalize(entry.get("name", "")) for entry in baseline}
    imported = [entry for entry in additions if normalize(entry["name"]) not in existing_names]
    root["version"] = "android-0.6"
    root["spells"] = baseline + imported
    catalog_path.write_text(json.dumps(root, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return len(imported)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--martial", type=Path, help="Path to 'Dубль, Мастера ближнего боя.docx'")
    parser.add_argument("--archmage", type=Path, help="Path to 'Книга Архимага.docx'")
    parser.add_argument(
        "--development-catalog",
        type=Path,
        default=Path("shared/src/commonMain/resources/fcp/dubl-3.69/content/development_catalog.json"),
    )
    parser.add_argument(
        "--magic-catalog",
        type=Path,
        default=Path("shared/src/commonMain/resources/fcp/dubl-3.69/content/magic_equipment_catalog.json"),
    )
    args = parser.parse_args()

    if not args.martial and not args.archmage:
        parser.error("At least one of --martial or --archmage is required")

    if args.martial:
        styles, techniques = parse_martial_arts(args.martial)
        if len(styles) != 19:
            raise SystemExit(f"Expected 19 approved unarmed/weapon styles, parsed {len(styles)}")
        if not techniques:
            raise SystemExit("No martial techniques parsed")
        update_development_catalog(args.development_catalog, styles, techniques)
        stance_count = sum("Стойка" in entry.get("tags", []) for entry in techniques)
        print(f"Imported {len(styles)} styles, {len(techniques)} unique techniques ({stance_count} stances).")

    if args.archmage:
        additions = parse_archmage_additions(args.archmage)
        if len(additions) != 59:
            raise SystemExit(f"Expected 59 additional spells from implemented schools, parsed {len(additions)}")
        imported = update_magic_catalog(args.magic_catalog, additions)
        print(f"Imported {imported} Archmage additions from {len(additions)} eligible definitions.")


if __name__ == "__main__":
    main()
