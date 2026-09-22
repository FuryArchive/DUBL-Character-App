import json
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from tools.rulebook.build_ruleset import build_ruleset


def test_build_ruleset_marks_bootstrap_mirror_and_surfaces_missing_and_ambiguous_sources(tmp_path: Path):
    source = tmp_path / "rules.docx"
    doc = Document()
    doc.add_heading("Навыки", 1)
    doc.add_heading("Alpha", 2)
    doc.add_paragraph("Alpha rule")
    doc.add_heading("Dup", 2)
    doc.add_paragraph("first")
    doc.add_heading("Dup", 2)
    doc.add_paragraph("second")
    doc.save(source)

    repo = tmp_path / "repo"
    shared = repo / "shared/src/commonMain/resources/fcp/dubl-3.69/content"
    shared.mkdir(parents=True)
    (shared / "development_catalog.json").write_text(json.dumps({"version":"x","entries":[
        {"id":"a","name":"Alpha"},{"id":"d","name":"Dup"},{"id":"m","name":"Missing"}
    ]}), encoding="utf-8")
    for name, root in {
        "chi_catalog.json":{"schools":[],"techniques":[]},
        "magic_equipment_catalog.json":{"spells":[],"gear":[]},
        "skill_effects_catalog.json":{"effects":[]},
    }.items():
        (shared / name).write_text(json.dumps(root), encoding="utf-8")

    out = tmp_path / "bundle"
    manifest = build_ruleset(source, repo, out)

    assert manifest["domains"]["development"]["status"] == "bootstrap_mirror"
    built = json.loads((out / "content/development.json").read_text(encoding="utf-8"))
    refs = {x["id"]: x.get("sourceRefs", []) for x in built["entries"]}
    assert refs["a"] == ["core:p-000002"]
    diagnostics = json.loads((out / "diagnostics.json").read_text(encoding="utf-8"))["diagnostics"]
    kinds = {(d["kind"], d["subject"]) for d in diagnostics}
    assert ("ambiguous-source", "development:d") in kinds
    assert ("missing-source", "development:m") in kinds

def test_build_ruleset_uses_explicit_source_policy_for_supplement_entities(tmp_path: Path):
    core = tmp_path / "core.docx"
    d = Document(); d.add_heading("Shared Name", 1); d.save(core)
    melee = tmp_path / "melee.docx"
    d = Document(); d.add_heading("Shared Name", 1); d.add_heading("Martial Move", 2); d.save(melee)
    archmage = tmp_path / "archmage.docx"
    d = Document(); d.add_heading("Arc Spell", 1); d.save(archmage)

    repo = tmp_path / "repo"; shared = repo / "shared/src/commonMain/resources/fcp/dubl-3.69/content"; shared.mkdir(parents=True)
    (shared / "development_catalog.json").write_text(json.dumps({"entries":[
        {"id":"feat_core","name":"Shared Name","section":"Навыки"},
        {"id":"martial_technique_x","name":"Martial Move","section":"Боевые искусства"}
    ]}), encoding="utf-8")
    (shared / "chi_catalog.json").write_text(json.dumps({"schools":[],"techniques":[]}), encoding="utf-8")
    (shared / "magic_equipment_catalog.json").write_text(json.dumps({"spells":[{"id":"archmage_x","name":"Arc Spell"}],"gear":[]}), encoding="utf-8")
    (shared / "skill_effects_catalog.json").write_text(json.dumps({"effects":[]}), encoding="utf-8")

    out = tmp_path / "bundle"
    build_ruleset({"core": core, "melee": melee, "archmage": archmage}, repo, out)
    dev = json.loads((out / "content/development.json").read_text(encoding="utf-8"))["entries"]
    magic = json.loads((out / "content/magic_equipment.json").read_text(encoding="utf-8"))["spells"]
    assert dev[0]["sourceRefs"] == ["core:p-000001"]
    assert dev[1]["sourceRefs"] == ["melee:p-000002"]
    assert magic[0]["sourceRefs"] == ["archmage:p-000001"]


def test_build_ruleset_prefers_explicit_archmage_source_metadata_over_legacy_id_prefix(tmp_path: Path):
    core = tmp_path / "core.docx"
    d = Document(); d.add_heading("Shared Spell", 1); d.save(core)
    archmage = tmp_path / "archmage.docx"
    d = Document(); d.add_heading("Shared Spell", 1); d.save(archmage)

    repo = tmp_path / "repo"
    shared = repo / "shared/src/commonMain/resources/fcp/dubl-3.69/content"
    shared.mkdir(parents=True)
    (shared / "development_catalog.json").write_text(json.dumps({"entries": []}), encoding="utf-8")
    (shared / "chi_catalog.json").write_text(json.dumps({"schools": [], "techniques": []}), encoding="utf-8")
    (shared / "magic_equipment_catalog.json").write_text(json.dumps({
        "spells": [{"id": "spell_without_legacy_prefix", "name": "Shared Spell", "source": "Книга Архимага"}],
        "gear": []
    }), encoding="utf-8")
    (shared / "skill_effects_catalog.json").write_text(json.dumps({"effects": []}), encoding="utf-8")

    out = tmp_path / "bundle"
    build_ruleset({"core": core, "archmage": archmage}, repo, out)
    spell = json.loads((out / "content/magic_equipment.json").read_text(encoding="utf-8"))["spells"][0]
    assert spell["sourcePolicy"] == ["archmage"]
    assert spell["sourceRefs"] == ["archmage:p-000001"]

def test_build_ruleset_can_link_unique_inline_rule_name_for_bootstrap_domain(tmp_path: Path):
    core = tmp_path / "core.docx"; d=Document(); d.add_heading("Root",1); d.add_paragraph("Разряд 2 Ци — электрический удар"); d.save(core)
    repo=tmp_path/"repo"; shared=repo/"shared/src/commonMain/resources/fcp/dubl-3.69/content"; shared.mkdir(parents=True)
    (shared/"development_catalog.json").write_text(json.dumps({"entries":[]}),encoding="utf-8")
    (shared/"chi_catalog.json").write_text(json.dumps({"schools":[],"techniques":[]}),encoding="utf-8")
    (shared/"magic_equipment_catalog.json").write_text(json.dumps({"spells":[],"gear":[]}),encoding="utf-8")
    (shared/"skill_effects_catalog.json").write_text(json.dumps({"effects":[{"id":"effect_x","name":"Разряд"}]}),encoding="utf-8")
    out=tmp_path/"bundle"; build_ruleset(core,repo,out)
    effect=json.loads((out/"content/skill_effects.json").read_text(encoding="utf-8"))["effects"][0]
    assert effect["sourceRefs"] == ["core:p-000002"]
    assert effect["sourceMatch"] == "inline-unique"

def test_build_ruleset_includes_source_diagnostics(tmp_path: Path):
    core=tmp_path/"core.docx"; d=Document(); d.add_heading("Draft",1); d.add_paragraph("Позже напишу ???"); d.save(core)
    repo=tmp_path/"repo"; shared=repo/"shared/src/commonMain/resources/fcp/dubl-3.69/content"; shared.mkdir(parents=True)
    (shared/"development_catalog.json").write_text(json.dumps({"entries":[]}),encoding="utf-8")
    (shared/"chi_catalog.json").write_text(json.dumps({"schools":[],"techniques":[]}),encoding="utf-8")
    (shared/"magic_equipment_catalog.json").write_text(json.dumps({"spells":[],"gear":[]}),encoding="utf-8")
    (shared/"skill_effects_catalog.json").write_text(json.dumps({"effects":[]}),encoding="utf-8")
    out=tmp_path/"bundle"; build_ruleset(core,repo,out)
    diags=json.loads((out/"diagnostics.json").read_text(encoding="utf-8"))["diagnostics"]
    assert any(d["kind"] == "draft-marker" for d in diags)

def test_build_ruleset_generates_conditions_from_core_source(tmp_path: Path):
    core=tmp_path/"core.docx"; d=Document(); d.add_heading("Состояния и эффекты",1); d.add_heading("Слепота",4); d.add_paragraph("Не видит."); d.add_heading("Типы урона",2); d.save(core)
    repo=tmp_path/"repo"; shared=repo/"shared/src/commonMain/resources/fcp/dubl-3.69/content"; shared.mkdir(parents=True)
    (shared/"development_catalog.json").write_text(json.dumps({"entries":[]}),encoding="utf-8")
    (shared/"chi_catalog.json").write_text(json.dumps({"schools":[],"techniques":[]}),encoding="utf-8")
    (shared/"magic_equipment_catalog.json").write_text(json.dumps({"spells":[],"gear":[]}),encoding="utf-8")
    (shared/"skill_effects_catalog.json").write_text(json.dumps({"effects":[]}),encoding="utf-8")
    out=tmp_path/"bundle"; manifest=build_ruleset(core,repo,out)
    assert manifest["domains"]["conditions"]["status"] == "source_generated"
    conditions=json.loads((out/"content/conditions.json").read_text(encoding="utf-8"))["conditions"]
    assert [x["name"] for x in conditions] == ["Слепота"]


def test_build_ruleset_reports_missing_source_generated_section_as_error_diagnostic(tmp_path: Path):
    core = tmp_path / "core.docx"
    d = Document()
    d.add_heading("Навыки", 1)
    d.add_paragraph("No conditions section here")
    d.save(core)

    repo = tmp_path / "repo"
    shared = repo / "shared/src/commonMain/resources/fcp/dubl-3.69/content"
    shared.mkdir(parents=True)
    (shared / "development_catalog.json").write_text(json.dumps({"entries": []}), encoding="utf-8")
    (shared / "chi_catalog.json").write_text(json.dumps({"schools": [], "techniques": []}), encoding="utf-8")
    (shared / "magic_equipment_catalog.json").write_text(json.dumps({"spells": [], "gear": []}), encoding="utf-8")
    (shared / "skill_effects_catalog.json").write_text(json.dumps({"effects": []}), encoding="utf-8")

    out = tmp_path / "bundle"
    manifest = build_ruleset(core, repo, out)

    conditions = json.loads((out / "content/conditions.json").read_text(encoding="utf-8"))
    assert conditions == {"schemaVersion": 1, "source": "core", "conditions": [], "mechanics": []}
    assert manifest["domains"]["conditions"]["coverage"]["missing"] == 1
    diagnostics = json.loads((out / "diagnostics.json").read_text(encoding="utf-8"))["diagnostics"]
    missing = [d for d in diagnostics if d["kind"] == "source-section-missing" and d["subject"] == "conditions"]
    assert len(missing) == 1
    assert missing[0]["severity"] == "error"


def test_build_ruleset_copies_tracked_explicit_resolutions_into_bundle(tmp_path: Path):
    core = tmp_path / "core.docx"
    d = Document(); d.add_heading("Состояния и эффекты", 1); d.add_heading("Слепота", 4); d.add_paragraph("Не видит."); d.add_heading("Типы урона", 2); d.save(core)

    repo = tmp_path / "repo"
    shared = repo / "shared/src/commonMain/resources/fcp/dubl-3.69/content"
    shared.mkdir(parents=True)
    (shared / "development_catalog.json").write_text(json.dumps({"entries": []}), encoding="utf-8")
    (shared / "chi_catalog.json").write_text(json.dumps({"schools": [], "techniques": []}), encoding="utf-8")
    (shared / "magic_equipment_catalog.json").write_text(json.dumps({"spells": [], "gear": []}), encoding="utf-8")
    (shared / "skill_effects_catalog.json").write_text(json.dumps({"effects": []}), encoding="utf-8")
    tracked = repo / "rulesets/dubl-3.69/resolutions.json"
    tracked.parent.mkdir(parents=True)
    payload = {"schemaVersion": 1, "ruleset": "dubl-3.69", "resolutions": [{
        "diagnosticId": "diag_example", "decision": "example", "rationale": "human decision"
    }]}
    tracked.write_text(json.dumps(payload), encoding="utf-8")

    out = tmp_path / "bundle"
    build_ruleset(core, repo, out)
    assert json.loads((out / "resolutions.json").read_text(encoding="utf-8")) == payload


def test_build_ruleset_carries_tracked_domain_runtime_artifact_into_manifest(tmp_path: Path):
    core = tmp_path / "core.docx"
    d = Document(); d.add_heading("Состояния и эффекты", 1); d.add_heading("Слепота", 4); d.add_paragraph("Не видит."); d.add_heading("Типы урона", 2); d.save(core)

    repo = tmp_path / "repo"
    shared = repo / "shared/src/commonMain/resources/fcp/dubl-3.69/content"
    shared.mkdir(parents=True)
    for name, payload in {
        "development_catalog.json": {"entries": []},
        "chi_catalog.json": {"schools": [], "techniques": []},
        "magic_equipment_catalog.json": {"spells": [], "gear": []},
        "skill_effects_catalog.json": {"effects": []},
    }.items():
        (shared / name).write_text(json.dumps(payload), encoding="utf-8")
    config = repo / "rulesets/dubl-3.69/config.json"
    config.parent.mkdir(parents=True)
    config.write_text(json.dumps({
        "domains": {
            "conditions": {
                "status": "source_generated",
                "sources": ["core"],
                "runtimeArtifact": "shared/src/commonMain/resources/fcp/dubl-3.69/content/conditions_catalog.json",
                "compiledArtifact": {
                    "type": "skill_catalog_kotlin",
                    "path": "shared/src/commonMain/kotlin/example.kt",
                },
            }
        }
    }), encoding="utf-8")

    out = tmp_path / "bundle"
    manifest = build_ruleset(core, repo, out)
    assert manifest["domains"]["conditions"]["runtimeArtifact"] == "shared/src/commonMain/resources/fcp/dubl-3.69/content/conditions_catalog.json"
    assert manifest["domains"]["conditions"]["compiledArtifact"] == {
        "type": "skill_catalog_kotlin",
        "path": "shared/src/commonMain/kotlin/example.kt",
    }


def test_build_ruleset_emits_source_generated_skill_candidate_from_core_table_and_bindings(tmp_path: Path):
    core = tmp_path / "core.docx"
    d = Document()
    d.add_heading("Умения", 1)
    table = d.add_table(rows=1, cols=5)
    for cell, value in zip(table.rows[0].cells, ["Умение", "Описание", "Авто 6", "Авто 12", "Используется нетренированным"]):
        cell.text = value
    row = table.add_row()
    for cell, value in zip(row.cells, ["Атлетика", "Прыжки", "Да", "Нет", "Да"]):
        cell.text = value
    costs = d.add_table(rows=1, cols=3)
    for cell, value in zip(costs.rows[0].cells, ["Значение", "Стоимость", "Стоимость для поднятия ранга с ноля"]):
        cell.text = value
    row = costs.add_row()
    for cell, value in zip(row.cells, ["0 → 1", "10", "10"]):
        cell.text = value
    d.add_heading("Состояния и эффекты", 1); d.add_heading("Слепота", 4); d.add_paragraph("Не видит."); d.add_heading("Типы урона", 2)
    d.save(core)

    repo = tmp_path / "repo"
    shared = repo / "shared/src/commonMain/resources/fcp/dubl-3.69/content"
    shared.mkdir(parents=True)
    for name, payload in {
        "development_catalog.json": {"entries": []},
        "chi_catalog.json": {"schools": [], "techniques": []},
        "magic_equipment_catalog.json": {"spells": [], "gear": []},
        "skill_effects_catalog.json": {"effects": []},
    }.items():
        (shared / name).write_text(json.dumps(payload), encoding="utf-8")
    bindings = repo / "rulesets/dubl-3.69/bindings/skills.json"
    bindings.parent.mkdir(parents=True)
    bindings.write_text(json.dumps({"bindings": [
        {"id": "athletics", "name": "Атлетика", "category": "PHYSICAL", "defaultAttributeHint": "STRENGTH", "template": False},
        {"id": "computers", "name": "Компьютеры", "category": "TECHNICAL", "defaultAttributeHint": "INTELLIGENCE", "template": False},
    ]}), encoding="utf-8")
    config = repo / "rulesets/dubl-3.69/config.json"
    config.parent.mkdir(parents=True, exist_ok=True)
    config.write_text(json.dumps({"domains": {"skills": {"status": "source_generated_candidate", "sources": ["core"]}}}), encoding="utf-8")

    out = tmp_path / "bundle"
    manifest = build_ruleset(core, repo, out)
    skills = json.loads((out / "content/skills.json").read_text(encoding="utf-8"))
    assert [item["id"] for item in skills["skills"]] == ["athletics"]
    assert skills["rankCosts"] == [0, 10]
    assert manifest["domains"]["skills"]["status"] == "source_generated_candidate"
    assert manifest["domains"]["skills"]["coverage"] == {"total": 2, "linked": 1, "ambiguous": 0, "missing": 1}
    diagnostics = json.loads((out / "diagnostics.json").read_text(encoding="utf-8"))["diagnostics"]
    assert any(d["kind"] == "skill-binding-not-in-base-table" and d["subject"] == "skills:computers" for d in diagnostics)


def test_build_ruleset_emits_regular_development_candidate_and_runtime_drift(tmp_path: Path):
    core = tmp_path / "core.docx"
    d = Document()
    d.add_heading("Навыки", 1)
    d.add_heading("Атлетика", 2)
    d.add_paragraph("Плавание")
    d.add_paragraph("Стоимость: 10")
    d.add_paragraph("Ранги: 2")
    d.add_paragraph("Требование: Ловкость 2")
    d.add_paragraph("Выгода: Канон")
    d.add_heading("Состояния и эффекты", 1); d.add_heading("Слепота", 4); d.add_paragraph("Не видит."); d.add_heading("Типы урона", 2)
    d.save(core)

    repo = tmp_path / "repo"
    shared = repo / "shared/src/commonMain/resources/fcp/dubl-3.69/content"
    shared.mkdir(parents=True)
    (shared / "development_catalog.json").write_text(json.dumps({"entries": [{
        "id": "feat-swim", "name": "Плавание", "section": "Навыки", "category": "Атлетика",
        "cost": 99, "ranks": 2, "requirements": "Ловкость 2", "benefit": "Канон"
    }]}), encoding="utf-8")
    for name, payload in {
        "chi_catalog.json": {"schools": [], "techniques": []},
        "magic_equipment_catalog.json": {"spells": [], "gear": []},
        "skill_effects_catalog.json": {"effects": []},
    }.items():
        (shared / name).write_text(json.dumps(payload), encoding="utf-8")
    bindings = repo / "rulesets/dubl-3.69/bindings/development_regular.json"
    bindings.parent.mkdir(parents=True)
    bindings.write_text(json.dumps({"bindings": [{"id": "feat-swim", "name": "Плавание", "category": "Атлетика"}]}), encoding="utf-8")

    out = tmp_path / "bundle"
    manifest = build_ruleset(core, repo, out)

    candidate = json.loads((out / "content/development_regular.json").read_text(encoding="utf-8"))
    assert candidate["entries"][0]["cost"] == 10
    assert manifest["domains"]["development_regular"]["status"] == "source_generated_candidate"
    assert manifest["domains"]["development_regular"]["coverage"] == {"total": 1, "linked": 1, "ambiguous": 0, "missing": 0}
    diagnostics = json.loads((out / "diagnostics.json").read_text(encoding="utf-8"))["diagnostics"]
    drift = [d for d in diagnostics if d["kind"] == "development-runtime-drift"]
    assert len(drift) == 1
    assert drift[0]["differences"] == {"cost": {"rulebook": 10, "runtime": 99}}
